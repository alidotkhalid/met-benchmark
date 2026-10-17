"""Score the abstention classifier against the human labels in abstention_annotation.xlsx.

The sample is stratified by model and by the classifier's own decision (50 readings it called a refusal and 50 it
called an answer per model, fewer where too few distinct readings exist), so raw counts over the sample are NOT
population rates. Within each stratum we estimate the share of true refusals, then weight by the stratum's size in
the full pool of isolated readings (stored in abstention_key.json):

  precision (refusal)  = share of predicted refusals that are true refusals        (weighted over models)
  recall    (refusal)  = true refusals caught / all true refusals                  (weighted)
  accuracy             = share of all readings classified correctly                 (weighted)

Gold label: the consensus where both annotators labelled an item and agree; items where they disagree are listed
for adjudication and, if an 'Adjudicated' column is added (G), that value is used. With one annotator only, that
annotator's labels are used. Cohen's kappa is reported whenever both annotators have labelled.

usage: python score_abstention.py [abstention_annotation.xlsx] [abstention_key.json]
"""
import json, sys
from collections import defaultdict
from openpyxl import load_workbook

xl = sys.argv[1] if len(sys.argv) > 1 else "abstention_annotation.xlsx"
kp = sys.argv[2] if len(sys.argv) > 2 else "abstention_key.json"
key = json.load(open(kp))
pred = {k["item"]: k for k in key["items"]}
sizes = key["strata_sizes"]

ws = load_workbook(xl, data_only=True)["Items"]
A, B, ADJ = {}, {}, {}
for row in ws.iter_rows(min_row=2, values_only=True):
    item = row[0]
    if not item:
        continue
    norm = lambda v: (str(v).strip().lower() if v not in (None, "") else None)
    A[item], B[item] = norm(row[3]), norm(row[4])
    if len(row) > 6:
        ADJ[item] = norm(row[6])

both = [i for i in A if A[i] and B.get(i)]
if both:
    agree = sum(A[i] == B[i] for i in both)
    po = agree / len(both)
    pa = sum(A[i] == "refusal" for i in both) / len(both)
    pb = sum(B[i] == "refusal" for i in both) / len(both)
    pe = pa * pb + (1 - pa) * (1 - pb)
    kappa = (po - pe) / (1 - pe) if pe < 1 else float("nan")
    print(f"inter-annotator: {len(both)} items labelled by both, agreement {100*po:.1f}%, Cohen's kappa {kappa:.3f}")
    dis = [i for i in both if A[i] != B[i] and not ADJ.get(i)]
    if dis:
        print(f"  {len(dis)} disagreements without adjudication (excluded from scoring):", ", ".join(sorted(dis)))

gold = {}
for i in A:
    if ADJ.get(i):
        gold[i] = ADJ[i]
    elif A[i] and B.get(i):
        if A[i] == B[i]:
            gold[i] = A[i]
    elif A[i] or B.get(i):
        gold[i] = A[i] or B[i]

cells = defaultdict(lambda: [0, 0])      # (tag, pred) -> [n labelled, n gold refusal]
for i, g in gold.items():
    k = (pred[i]["tag"], pred[i]["pred"])
    cells[k][0] += 1
    cells[k][1] += g == "refusal"
print(f"\nscored items: {len(gold)} of {len(pred)}")

tp = fp = fn = tn = 0.0
for tag in sorted({t for t, _ in cells}):
    for p in ("refusal", "answer"):
        n, r = cells.get((tag, p), [0, 0])
        if not n:
            continue
        N = sizes[f"{tag}|{p}"]
        share = r / n                    # share of true refusals in this stratum
        print(f"  {tag} predicted {p:8} labelled {n:3}  true refusals {r:3} ({100*share:5.1f}%)  stratum size {N}")
        if p == "refusal":
            tp += N * share; fp += N * (1 - share)
        else:
            fn += N * share; tn += N * (1 - share)
tot = tp + fp + fn + tn
if tot:
    prec = tp / (tp + fp) if tp + fp else float("nan")
    rec = tp / (tp + fn) if tp + fn else float("nan")
    print(f"\nweighted over the full pool of isolated readings:")
    print(f"  refusal precision {100*prec:.1f}%   refusal recall {100*rec:.1f}%   accuracy {100*(tp+tn)/tot:.1f}%")
    print(f"  estimated true refusal rate {100*(tp+fn)/tot:.1f}%   classifier refusal rate {100*(tp+fp)/tot:.1f}%")
